from __future__ import annotations


import os
import subprocess
import sys
from pathlib import Path


# 換到其他電腦時，只需修改下一行的 pythonw.exe 絕對路徑。
PYTHONW_EXE = Path(
    r"C:\Users\python3.8\win64\pythonw.exe"
)


def _executable_path_key(path: Path) -> str:
    """Return a case-insensitive absolute identity for a Windows executable."""

    return os.path.normcase(os.path.abspath(str(path)))


def _show_bootstrap_error(detail: str) -> None:
    """Show an error even when the current interpreter has no console."""

    if sys.stderr is not None:
        sys.stderr.write(detail + "\n")
    try:
        import ctypes

        ctypes.windll.user32.MessageBoxW(0, detail, "TSM 圖片批次匯入 Excel", 0x10)
    except (AttributeError, OSError):
        pass


def _ensure_configured_python() -> None:
    """Relaunch this file with the configured portable Python when necessary."""

    configured_pythonw = PYTHONW_EXE.expanduser()
    configured_python = configured_pythonw.with_name("python.exe")
    current_key = _executable_path_key(Path(sys.executable))
    configured_keys = {
        _executable_path_key(configured_pythonw),
        _executable_path_key(configured_python),
    }
    if current_key in configured_keys:
        return

    if not configured_pythonw.is_file():
        _show_bootstrap_error(
            "找不到指定的 Python 執行檔：\n{}\n\n"
            "請用文字編輯器開啟 tsm_image_excel_gui.pyw，修改檔案最上方的 PYTHONW_EXE。".format(configured_pythonw)
        )
        raise SystemExit(1)

    current_is_console_python = Path(sys.executable).name.lower() == "python.exe"
    target_executable = (
        configured_python if current_is_console_python and configured_python.is_file() else configured_pythonw
    )
    script_path = Path(__file__).resolve()
    command = [str(target_executable), str(script_path)] + sys.argv[1:]
    try:
        if current_is_console_python:
            return_code = subprocess.call(command, cwd=str(script_path.parent))
            raise SystemExit(return_code)
        subprocess.Popen(command, cwd=str(script_path.parent))
    except OSError as exc:
        _show_bootstrap_error("無法使用指定的 Python 啟動程式：\n{}\n\n{}".format(target_executable, exc))
        raise SystemExit(1)
    raise SystemExit(0)


if __name__ == "__main__":
    _ensure_configured_python()


import re
import traceback
from dataclasses import dataclass
from typing import Iterable

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.drawing.image import Image as ExcelImage
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit(
        "Missing dependency. Please install with: pip install openpyxl pillow"
    ) from exc


IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".gif", ".tif", ".tiff"}
NUMBER_RE = re.compile(r"(\d+)")
FILTER_CONDITION_COUNT = 4


@dataclass(frozen=True)
class ImageEntry:
    path: Path
    sort_key: tuple


def natural_key(value: str) -> tuple:
    parts = NUMBER_RE.split(value)
    key: list[object] = []
    for part in parts:
        key.append(int(part) if part.isdigit() else part.lower())
    return tuple(key)


def excel_width_to_pixels(width: float) -> int:
    return max(1, int(width * 7 + 5))


def points_to_pixels(points: float) -> int:
    return max(1, int(points * 96 / 72))


def pixels_to_points(pixels: float) -> float:
    return max(1.0, pixels * 72 / 96)


def compute_scaled_size(native_w: int, native_h: int, box_w: int, box_h: int) -> tuple[int, int]:
    """Shrink-only, aspect-ratio-preserving fit of (native_w, native_h) into (box_w, box_h)."""
    if native_w <= 0 or native_h <= 0:
        return box_w, box_h
    scale = min(box_w / native_w, box_h / native_h, 1.0)
    return max(1, round(native_w * scale)), max(1, round(native_h * scale))


def safe_sheet_name(name: str, used_names: set[str]) -> str:
    invalid = '[]:*?/\\'
    cleaned = "".join("_" if char in invalid else char for char in name).strip()
    cleaned = cleaned or "Sheet"
    cleaned = cleaned[:31]
    candidate = cleaned
    index = 2
    while candidate in used_names:
        suffix = f"_{index}"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        index += 1
    used_names.add(candidate)
    return candidate


def read_temperature_list(path: Path) -> list:
    """Read column A from row 1 downward until the first empty cell."""
    workbook = load_workbook(path, data_only=True)
    try:
        sheet = workbook.active
        values: list = []
        row = 1
        while True:
            cell_value = sheet.cell(row=row, column=1).value
            if cell_value is None or (isinstance(cell_value, str) and not cell_value.strip()):
                break
            values.append(cell_value)
            row += 1
        return values
    finally:
        workbook.close()


def matches_conditions(name: str, conditions: list[str]) -> bool:
    if not conditions:
        return True
    return all(keyword in name for keyword in conditions)


def iter_image_entries(folder: Path, conditions: list[str]) -> list[ImageEntry]:
    entries: list[ImageEntry] = []
    for path in folder.iterdir():
        if not (path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS):
            continue
        if not matches_conditions(path.name, conditions):
            continue
        entries.append(ImageEntry(path=path, sort_key=natural_key(path.stem)))
    return sorted(entries, key=lambda item: item.sort_key)


def iter_directories(root: Path) -> list[Path]:
    """Immediate (one level) subfolders of root, sorted naturally by name."""
    return sorted(
        (path for path in root.iterdir() if path.is_dir()),
        key=lambda path: natural_key(path.name),
    )


def group_subfolder_sheets(
    rows: list[tuple[str, Path]], log=None
) -> tuple[list[tuple[str, list[tuple[str, Path]]]], list[str]]:
    """For each (name, path), search one level deep for subfolders. Each row
    becomes one worksheet (named `name`); its subfolders become the columns
    (one leg per subfolder, named after the subfolder itself) inside that sheet."""
    def _log(msg: str) -> None:
        if log:
            log(msg)

    groups: list[tuple[str, list[tuple[str, Path]]]] = []
    warnings: list[str] = []
    for name, path in rows:
        _log(f"── 掃描上層資料夾: 「{name}」  ({path})")
        if not path.exists() or not path.is_dir():
            msg = f"路徑不存在: {path}"
            _log(f"   ✗ {msg}")
            warnings.append(msg)
            continue
        subdirs = iter_directories(path)
        if not subdirs:
            msg = f"「{name}」底下沒有子資料夾（限定搜尋一層，請確認路徑）"
            _log(f"   ⚠ {msg}")
            warnings.append(msg)
            continue
        _log(f"   找到 {len(subdirs)} 個子資料夾: {[d.name for d in subdirs]}")
        columns = [(subdir.name, subdir) for subdir in subdirs]
        groups.append((name, columns))
    return groups, warnings


def populate_sheet(
    sheet,
    legs: list[tuple[str, Path]],
    column_width: float,
    image_row_height: float,
    temperatures: list,
    conditions: list[str],
    header_row_height: float = 20.0,
    log=None,
) -> tuple[int, list[str], list[str]]:
    """Fill a worksheet with one column per (name, path) leg. Returns
    (total_images, warnings, mismatch_warnings)."""
    def _log(msg: str) -> None:
        if log:
            log(msg)

    sheet.freeze_panes = "A1"
    sheet.sheet_view.showGridLines = True

    box_w = excel_width_to_pixels(column_width)
    box_h = points_to_pixels(image_row_height)

    warnings: list[str] = []
    mismatch_warnings: list[str] = []
    total_images = 0
    columns_data: list[dict] = []

    for column_index, (name, path) in enumerate(legs, start=1):
        _log(f"── 處理欄: 「{name}」  ({path})")
        column_letter = get_column_letter(column_index)
        sheet.column_dimensions[column_letter].width = column_width

        header_cell = sheet.cell(row=1, column=column_index, value=name)
        header_cell.alignment = Alignment(horizontal="center", vertical="center")
        header_cell.font = Font(bold=True)
        sheet.row_dimensions[1].height = header_row_height

        if not path.exists() or not path.is_dir():
            msg = f"路徑不存在，略過: {path}"
            _log(f"   ✗ {msg}")
            warnings.append(msg)
            columns_data.append({"placements": []})
            continue

        images = iter_image_entries(path, conditions)
        _log(f"   {len(images)} 張符合圖片")
        if not images:
            msg = f"{name}: 無符合圖片。"
            _log(f"   ⚠ {msg}")
            warnings.append(msg)

        if temperatures and len(images) != len(temperatures):
            msg = f"「{name}」符合條件的照片數量({len(images)}) 與溫度表數量({len(temperatures)}) 不一致"
            _log(f"   ⚠ {msg}")
            warnings.append(msg)
            mismatch_warnings.append(msg)

        placements = []
        for entry in images:
            try:
                excel_image = ExcelImage(str(entry.path))
            except Exception as img_exc:
                import traceback as _tb
                msg = f"圖片載入失敗: {entry.path.name} [{type(img_exc).__name__}] {img_exc}"
                _log(f"     X {msg}")
                _log(f"       {_tb.format_exc().splitlines()[-1]}")
                warnings.append(msg)
                continue
            native_w, native_h = excel_image.width, excel_image.height
            final_w, final_h = compute_scaled_size(native_w, native_h, box_w, box_h)
            placements.append((entry, excel_image, final_w, final_h))

        columns_data.append({"placements": placements})

    max_blocks = max((len(col["placements"]) for col in columns_data), default=0)
    row_height_px = [0] * max_blocks
    for col in columns_data:
        for i, (_entry, _img, _w, h) in enumerate(col["placements"]):
            row_height_px[i] = max(row_height_px[i], h)

    for column_index, col in enumerate(columns_data, start=1):
        for serial, (entry, excel_image, final_w, final_h) in enumerate(col["placements"], start=1):
            base_row = (serial - 1) * 3 + 2
            temp_value = temperatures[serial - 1] if serial - 1 < len(temperatures) else "(無對應溫度)"

            serial_cell = sheet.cell(row=base_row, column=column_index, value=serial)
            temp_cell = sheet.cell(row=base_row + 1, column=column_index, value=temp_value)
            image_cell = sheet.cell(row=base_row + 2, column=column_index)

            serial_cell.alignment = Alignment(horizontal="center", vertical="center")
            temp_cell.alignment = Alignment(horizontal="center", vertical="center")
            serial_cell.font = Font(bold=True)

            sheet.row_dimensions[base_row].height = header_row_height
            sheet.row_dimensions[base_row + 1].height = header_row_height
            sheet.row_dimensions[base_row + 2].height = pixels_to_points(row_height_px[serial - 1])

            excel_image.width = final_w
            excel_image.height = final_h
            sheet.add_image(excel_image, image_cell.coordinate)
            total_images += 1
            _log(f"   [{entry.path.parent.name}] [{serial}] {entry.path.name} → 溫度={temp_value}")

    return total_images, warnings, mismatch_warnings


def scan_legs(
    legs: list[tuple[str, Path]],
    temperatures: list,
    conditions: list[str],
    log=None,
) -> tuple[list[str], list[str]]:
    """Dry-run: scan folders and log findings without writing anything."""
    def _log(msg: str) -> None:
        if log:
            log(msg)

    warnings: list[str] = []
    mismatch_warnings: list[str] = []
    for name, path in legs:
        _log(f"── 欄: 「{name}」  ({path})")
        if not path.exists() or not path.is_dir():
            msg = f"路徑不存在: {path}"
            _log(f"   ✗ {msg}")
            warnings.append(msg)
            continue
        images = iter_image_entries(path, conditions)
        _log(f"   {len(images)} 張符合圖片")
        for entry in images:
            _log(f"      {entry.path.name}")
        if not images:
            warnings.append(f"{name}: 無符合圖片")
        if temperatures and len(images) != len(temperatures):
            msg = f"「{name}」符合條件的照片數量({len(images)}) 與溫度表數量({len(temperatures)}) 不一致"
            _log(f"   ⚠ {msg}")
            warnings.append(msg)
            mismatch_warnings.append(msg)
    return warnings, mismatch_warnings


def build_workbook(
    column_legs: list[tuple[str, Path]],
    subfolder_rows: list[tuple[str, Path]],
    output_path: Path,
    column_width: float,
    image_row_height: float,
    temperatures: list,
    conditions: list[str],
    header_row_height: float = 20.0,
    log=None,
) -> tuple[int, int, list[str], list[str]]:
    def _log(msg: str) -> None:
        if log:
            log(msg)

    workbook = Workbook()
    used_sheet_names: set[str] = set()
    warnings: list[str] = []
    mismatch_warnings: list[str] = []
    total_images = 0
    sheet_count = 0

    if column_legs:
        sheet = workbook.active
        sheet.title = safe_sheet_name(output_path.stem, used_sheet_names)
        _log(f"── 建立工作表(依欄輸出): 「{sheet.title}」")
        imgs, w, m = populate_sheet(
            sheet, column_legs, column_width, image_row_height,
            temperatures, conditions, header_row_height, log,
        )
        total_images += imgs
        warnings += w
        mismatch_warnings += m
        sheet_count += 1
    else:
        workbook.remove(workbook.active)

    groups, expand_warnings = group_subfolder_sheets(subfolder_rows, log=log)
    warnings += expand_warnings
    for name, columns in groups:
        sheet = workbook.create_sheet(safe_sheet_name(name, used_sheet_names))
        _log(f"── 建立工作表(子資料夾分頁): 「{sheet.title}」")
        imgs, w, m = populate_sheet(
            sheet, columns, column_width, image_row_height,
            temperatures, conditions, header_row_height, log,
        )
        total_images += imgs
        warnings += w
        mismatch_warnings += m
        sheet_count += 1

    _log(f"儲存 Excel → {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    _log("儲存完成。")
    return sheet_count, total_images, warnings, mismatch_warnings


class LegRow:
    def __init__(self, parent: ttk.Frame, on_remove) -> None:
        self.frame = ttk.Frame(parent)
        self.name_var = tk.StringVar()
        self.path_var = tk.StringVar()

        ttk.Label(self.frame, text="資料夾名稱").pack(side="left", padx=(0, 2))
        ttk.Entry(self.frame, textvariable=self.name_var, width=10).pack(side="left", padx=(0, 6))
        ttk.Label(self.frame, text="路徑").pack(side="left", padx=(0, 2))
        ttk.Entry(self.frame, textvariable=self.path_var).pack(
            side="left", fill="x", expand=True, padx=(0, 4)
        )
        ttk.Button(self.frame, text="瀏覽", command=self._browse, width=5).pack(
            side="left", padx=(0, 4)
        )
        ttk.Button(self.frame, text="×", command=lambda: on_remove(self), width=2).pack(
            side="left"
        )

        self.path_var.trace_add("write", self._on_path_changed)

    def _on_path_changed(self, *_args) -> None:
        """Keep 資料夾名稱 in sync when the path is typed or pasted (e.g. Ctrl+V)."""
        raw = self.path_var.get().strip()
        if not raw:
            return
        name = Path(raw.strip('"\'')).name.strip()
        if name:
            self.name_var.set(name)

    def _browse(self) -> None:
        folder = filedialog.askdirectory(title="選擇照片資料夾")
        if not folder:
            return
        self.path_var.set(folder)
        self.name_var.set(Path(folder).name)

    def pack(self, **kwargs) -> None:
        self.frame.pack(**kwargs)

    def destroy(self) -> None:
        self.frame.destroy()

    def get(self) -> tuple[str, str]:
        return self.name_var.get().strip(), self.path_var.get().strip()


class LegListPanel:
    """A labeled, scrollable list of LegRow entries with a row-count control."""

    def __init__(self, parent: ttk.Frame, title: str) -> None:
        self.rows: list[LegRow] = []
        self.add_count_var = tk.StringVar(value="1")

        container = ttk.Frame(parent)
        container.pack(fill=tk.BOTH, expand=True)
        container.columnconfigure(0, weight=1)
        container.rowconfigure(1, weight=1)

        header = ttk.Frame(container)
        header.grid(row=0, column=0, sticky="ew", pady=(0, 2))
        header.columnconfigure(0, weight=1)
        ttk.Label(header, text=title).grid(row=0, column=0, sticky="w")
        ttk.Label(header, text="欄位數量").grid(row=0, column=1, sticky="e", padx=(0, 4))
        ttk.Entry(header, textvariable=self.add_count_var, width=4).grid(
            row=0, column=2, sticky="e", padx=(0, 4)
        )
        ttk.Button(header, text="設定", command=self._set_count).grid(
            row=0, column=3, sticky="e"
        )

        list_container = ttk.LabelFrame(container, padding=4)
        list_container.grid(row=1, column=0, sticky="nsew")
        list_container.columnconfigure(0, weight=1)
        list_container.rowconfigure(0, weight=1)

        self._canvas = tk.Canvas(list_container, height=130, highlightthickness=0)
        scroll_y = ttk.Scrollbar(list_container, orient="vertical", command=self._canvas.yview)
        self._leg_frame = ttk.Frame(self._canvas)
        self._leg_frame.bind(
            "<Configure>",
            lambda e: self._canvas.configure(scrollregion=self._canvas.bbox("all")),
        )
        self._win_id = self._canvas.create_window((0, 0), window=self._leg_frame, anchor="nw")
        self._canvas.configure(yscrollcommand=scroll_y.set)
        self._canvas.bind(
            "<Configure>", lambda e: self._canvas.itemconfig(self._win_id, width=e.width)
        )

        self._canvas.grid(row=0, column=0, sticky="nsew")
        scroll_y.grid(row=0, column=1, sticky="ns")

        self._add_row()

    def _add_row(self) -> None:
        row = LegRow(self._leg_frame, on_remove=self._remove_row)
        row.pack(fill=tk.X, pady=2, padx=2)
        self.rows.append(row)

    def _remove_row(self, row: LegRow) -> None:
        if len(self.rows) <= 1:
            messagebox.showinfo("提示", "至少需要保留一列。")
            return
        self.rows.remove(row)
        row.destroy()

    def _set_count(self) -> None:
        try:
            count = int(self.add_count_var.get())
            assert count > 0
        except Exception:
            messagebox.showerror("錯誤", "數量必須是正整數。")
            return
        current = len(self.rows)
        if count > current:
            for _ in range(count - current):
                self._add_row()
        elif count < current:
            for row in self.rows[count:]:
                row.destroy()
            self.rows = self.rows[:count]

    def get_entries(self) -> list[tuple[str, str]]:
        return [row.get() for row in self.rows]


class ImageExcelApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("TSM 圖片批次匯入 Excel")
        self.geometry("840x720")
        self.minsize(720, 600)

        self.output_var = tk.StringVar()
        self.temp_path_var = tk.StringVar()
        self.column_width_var = tk.StringVar(value="28")
        self.image_row_height_var = tk.StringVar(value="150")
        self.status_var = tk.StringVar(value="請新增資料夾後開始。")

        self.filter_conditions: list[tuple[tk.BooleanVar, tk.StringVar]] = [
            (tk.BooleanVar(value=False), tk.StringVar(value=""))
            for _ in range(FILTER_CONDITION_COUNT)
        ]

        self._build_ui()

    def _build_ui(self) -> None:
        main = ttk.Frame(self, padding=14)
        main.pack(fill=tk.BOTH, expand=True)
        main.columnconfigure(0, weight=1)
        main.rowconfigure(0, weight=1)
        main.rowconfigure(5, weight=1)

        # --- Tabs ---
        notebook = ttk.Notebook(main)
        notebook.grid(row=0, column=0, sticky="nsew", pady=(0, 8))

        tab_columns = ttk.Frame(notebook, padding=8)
        tab_subfolders = ttk.Frame(notebook, padding=8)
        notebook.add(tab_columns, text="依欄輸出（同一張工作表）")
        notebook.add(tab_subfolders, text="依子資料夾分頁輸出（每列一個工作表）")

        self.column_panel = LegListPanel(
            tab_columns, "照片資料夾列表（每列 = 工作表中的一欄）"
        )
        self.subfolder_panel = LegListPanel(
            tab_subfolders,
            "上層資料夾列表（每列 = 一個工作表，工作表名稱＝資料夾名稱；"
            "往下搜尋一層子資料夾，每個子資料夾 = 該工作表裡的一欄）",
        )

        # --- Output path ---
        out_row = ttk.Frame(main)
        out_row.grid(row=1, column=0, sticky="ew", pady=(0, 8))
        out_row.columnconfigure(1, weight=1)
        ttk.Label(out_row, text="輸出 Excel").grid(row=0, column=0, sticky="w", padx=(0, 8))
        ttk.Entry(out_row, textvariable=self.output_var).grid(
            row=0, column=1, sticky="ew", padx=(0, 8)
        )
        ttk.Button(out_row, text="另存", command=self._choose_output).grid(row=0, column=2)

        # --- Temperature table ---
        temp_row = ttk.Frame(main)
        temp_row.grid(row=2, column=0, sticky="ew", pady=(0, 8))
        temp_row.columnconfigure(1, weight=1)
        ttk.Label(temp_row, text="溫度表(Excel)").grid(row=0, column=0, sticky="w", padx=(0, 8))
        ttk.Entry(temp_row, textvariable=self.temp_path_var).grid(
            row=0, column=1, sticky="ew", padx=(0, 8)
        )
        ttk.Button(temp_row, text="瀏覽", command=self._choose_temp_file).grid(row=0, column=2)
        ttk.Label(
            temp_row,
            text="（從 A1 往下讀到第一個空白儲存格為止，依序對應每一欄/每個分頁的每張照片）",
            foreground="gray",
        ).grid(row=1, column=0, columnspan=3, sticky="w", pady=(2, 0))

        # --- Settings ---
        settings = ttk.LabelFrame(main, text="Excel 設定", padding=10)
        settings.grid(row=3, column=0, sticky="ew", pady=(0, 8))
        settings.columnconfigure(1, weight=1)
        settings.columnconfigure(3, weight=1)

        ttk.Label(settings, text="欄寬").grid(row=0, column=0, sticky="w")
        ttk.Entry(settings, textvariable=self.column_width_var, width=8).grid(
            row=0, column=1, sticky="w", padx=(6, 20)
        )
        ttk.Label(settings, text="圖片列高上限").grid(row=0, column=2, sticky="w")
        ttk.Entry(settings, textvariable=self.image_row_height_var, width=8).grid(
            row=0, column=3, sticky="w", padx=6
        )
        ttk.Label(
            settings,
            text="圖片會依原始比例縮小以符合欄寬 × 列高上限，不會放大、不會變形。",
            foreground="gray",
        ).grid(row=1, column=0, columnspan=4, sticky="w", pady=(4, 0))

        # --- Filter conditions ---
        filters = ttk.LabelFrame(main, text="篩選條件（勾選的條件須同時符合，皆不勾 = 不篩選）", padding=10)
        filters.grid(row=4, column=0, sticky="ew", pady=(0, 8))
        for i, (enabled_var, text_var) in enumerate(self.filter_conditions):
            ttk.Checkbutton(filters, text=f"條件{i + 1}", variable=enabled_var).grid(
                row=i, column=0, sticky="w", padx=(0, 6), pady=2
            )
            ttk.Entry(filters, textvariable=text_var, width=24).grid(
                row=i, column=1, sticky="w", pady=2
            )
        filters.columnconfigure(1, weight=1)

        # --- Log ---
        self.log_text = tk.Text(main, height=7, wrap="word")
        self.log_text.grid(row=5, column=0, sticky="nsew", pady=(0, 6))

        # --- Button row ---
        btn_row = ttk.Frame(main)
        btn_row.grid(row=6, column=0, sticky="ew")
        btn_row.columnconfigure(0, weight=1)
        ttk.Label(btn_row, textvariable=self.status_var).grid(row=0, column=0, sticky="w")
        ttk.Button(btn_row, text="預覽", command=self._run_preview).grid(
            row=0, column=1, padx=(0, 6)
        )
        ttk.Button(btn_row, text="溫度表預覽", command=self._run_temp_preview).grid(
            row=0, column=2, padx=(0, 6)
        )
        ttk.Button(btn_row, text="開始產生 Excel", command=self._run_export).grid(
            row=0, column=3
        )

    def _choose_output(self) -> None:
        path = filedialog.asksaveasfilename(
            title="儲存 Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel workbook", "*.xlsx")],
        )
        if path:
            self.output_var.set(path)

    def _choose_temp_file(self) -> None:
        path = filedialog.askopenfilename(
            title="選擇溫度表 Excel",
            filetypes=[("Excel workbook", "*.xlsx")],
        )
        if path:
            self.temp_path_var.set(path)

    def append_log(self, text: str) -> None:
        self.log_text.insert(tk.END, text + "\n")
        self.log_text.see(tk.END)
        self.update_idletasks()

    def _validate_rows(
        self, entries: list[tuple[str, str]], label: str
    ) -> list[tuple[str, Path]] | None:
        legs: list[tuple[str, Path]] = []
        for i, (name, path_str) in enumerate(entries, start=1):
            if not name and not path_str:
                continue
            if not name:
                messagebox.showerror("錯誤", f"「{label}」第 {i} 列的資料夾名稱不可空白。")
                return None
            if not path_str:
                messagebox.showerror("錯誤", f"「{label}」第 {i} 列「{name}」的路徑不可空白。")
                return None
            legs.append((name, Path(path_str)))
        return legs

    def _parse_inputs(
        self,
    ) -> tuple[
        list[tuple[str, Path]], list[tuple[str, Path]], list, list[str], float, float
    ] | None:
        column_legs = self._validate_rows(self.column_panel.get_entries(), "依欄輸出")
        if column_legs is None:
            return None
        subfolder_rows = self._validate_rows(self.subfolder_panel.get_entries(), "子資料夾分頁")
        if subfolder_rows is None:
            return None
        if not column_legs and not subfolder_rows:
            messagebox.showerror("錯誤", "請至少在其中一個分頁新增照片資料夾。")
            return None

        temp_path_str = self.temp_path_var.get().strip()
        if not temp_path_str:
            messagebox.showerror("錯誤", "請先選擇溫度表 Excel 檔案。")
            return None
        temp_path = Path(temp_path_str)
        if not temp_path.exists():
            messagebox.showerror("錯誤", f"溫度表檔案不存在: {temp_path}")
            return None
        try:
            temperatures = read_temperature_list(temp_path)
        except Exception as exc:
            messagebox.showerror("錯誤", f"讀取溫度表失敗: {exc}")
            return None
        if not temperatures:
            messagebox.showerror("錯誤", "溫度表 A 欄沒有讀到任何溫度值。")
            return None

        conditions: list[str] = []
        for i, (enabled_var, text_var) in enumerate(self.filter_conditions, start=1):
            if not enabled_var.get():
                continue
            keyword = text_var.get().strip()
            if not keyword:
                messagebox.showerror("錯誤", f"條件{i} 已勾選但未輸入關鍵字。")
                return None
            conditions.append(keyword)

        try:
            column_width = float(self.column_width_var.get())
            assert column_width > 0
        except Exception:
            messagebox.showerror("錯誤", "欄寬必須是正數。")
            return None

        try:
            image_row_height = float(self.image_row_height_var.get())
            assert image_row_height > 0
        except Exception:
            messagebox.showerror("錯誤", "圖片列高上限必須是正數。")
            return None

        return column_legs, subfolder_rows, temperatures, conditions, column_width, image_row_height

    def _run_temp_preview(self) -> None:
        temp_path_str = self.temp_path_var.get().strip()
        if not temp_path_str:
            messagebox.showerror("錯誤", "請先選擇溫度表 Excel 檔案。")
            return
        temp_path = Path(temp_path_str)
        if not temp_path.exists():
            messagebox.showerror("錯誤", f"溫度表檔案不存在: {temp_path}")
            return
        try:
            temperatures = read_temperature_list(temp_path)
        except Exception as exc:
            messagebox.showerror("錯誤", f"讀取溫度表失敗: {exc}")
            return
        self.append_log(f"=== 溫度表預覽 ({temp_path.name}) ===")
        self.append_log(f"共讀取到 {len(temperatures)} 個溫度點: {temperatures}")
        self.status_var.set(f"溫度表共 {len(temperatures)} 個溫度點。")

    def _run_preview(self) -> None:
        self.log_text.delete("1.0", tk.END)
        result = self._parse_inputs()
        if result is None:
            return
        column_legs, subfolder_rows, temperatures, conditions, _, _ = result
        self.append_log("=== 預覽掃描 ===")
        self.append_log(f"溫度表共 {len(temperatures)} 筆: {temperatures}")
        if conditions:
            self.append_log(f"篩選條件(AND): {conditions}")
        try:
            all_warnings: list[str] = []
            all_mismatches: list[str] = []

            if column_legs:
                self.append_log("--- 依欄輸出 ---")
                w, m = scan_legs(column_legs, temperatures, conditions, log=self.append_log)
                all_warnings += w
                all_mismatches += m

            if subfolder_rows:
                self.append_log("--- 子資料夾分頁輸出 ---")
                groups, expand_warnings = group_subfolder_sheets(
                    subfolder_rows, log=self.append_log
                )
                all_warnings += expand_warnings
                for sheet_name, columns in groups:
                    self.append_log(f"── 工作表:「{sheet_name}」")
                    w2, m2 = scan_legs(columns, temperatures, conditions, log=self.append_log)
                    all_warnings += w2
                    all_mismatches += m2

            if all_warnings:
                self.append_log("─── 注意 ───")
                for w in all_warnings:
                    self.append_log(f"⚠ {w}")
            self.append_log("=== 預覽完成，確認無誤後按「開始產生 Excel」 ===")
            self.status_var.set("預覽完成。")
            if all_mismatches:
                messagebox.showwarning("照片數量與溫度表不符", "\n".join(all_mismatches))
        except Exception as exc:
            self.append_log(f"✗ {exc}")
            self.append_log(traceback.format_exc())
            self.status_var.set("預覽失敗。")

    def _run_export(self) -> None:
        self.log_text.delete("1.0", tk.END)
        result = self._parse_inputs()
        if result is None:
            return
        column_legs, subfolder_rows, temperatures, conditions, column_width, image_row_height = result

        output_text = self.output_var.get().strip()
        if not output_text:
            messagebox.showerror("錯誤", "請先設定輸出 Excel 路徑。")
            return
        output = Path(output_text)

        if output.exists():
            if not messagebox.askyesno("覆寫確認", f"檔案已存在，確定要覆寫嗎？\n{output}"):
                self.status_var.set("已取消。")
                return

        self.status_var.set("產生中...")
        try:
            sheet_count, image_count, warnings, mismatches = build_workbook(
                column_legs=column_legs,
                subfolder_rows=subfolder_rows,
                output_path=output,
                column_width=column_width,
                image_row_height=image_row_height,
                temperatures=temperatures,
                conditions=conditions,
                log=self.append_log,
            )
            self.append_log(f"完成: {sheet_count} 個工作表, {image_count} 張圖片。")
            for w in warnings:
                self.append_log(f"⚠ {w}")
            self.status_var.set("完成")
            if mismatches:
                messagebox.showwarning("照片數量與溫度表不符", "\n".join(mismatches))
            messagebox.showinfo("完成", f"Excel 已產生:\n{output}")
        except Exception as exc:
            self.status_var.set("失敗")
            self.append_log(f"✗ {str(exc)}")
            self.append_log(traceback.format_exc())
            messagebox.showerror("錯誤", str(exc))


def main(argv: Iterable[str] | None = None) -> int:
    app = ImageExcelApp()
    app.mainloop()
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
